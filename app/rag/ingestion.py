"""
Document ingestion pipeline.

Loading priority for PDFs:
  1. PyMuPDF (fitz) — fast path for mostly text PDFs
  2. unstructured hi_res — richer extraction for OCR / tables / images
  3. pypdf — fallback if pymupdf unavailable

Other formats use dedicated loaders from langchain-community.
"""
import base64
import csv
import mimetypes
import os
import re
import shutil
import uuid
import zipfile
from datetime import date, datetime
from dataclasses import dataclass, field
from pathlib import Path

from dotenv import load_dotenv
from langchain_core.documents import Document
from langchain_core.messages import HumanMessage
from langchain_text_splitters import RecursiveCharacterTextSplitter
from loguru import logger

from app.config import get_settings
from app.rag.vectorstore import (
    add_raw_documents,
    add_summary_documents,
    delete_document_family,
)

settings = get_settings()
_ENV_FILE = Path(__file__).resolve().parent.parent.parent / ".env"

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".md", ".xlsx", ".xls", ".csv"}


@dataclass
class IngestionBundle:
    raw_docs: list[Document]
    summary_docs: list[Document] = field(default_factory=list)
    used_multimodal: bool = False


# ── Groq helpers ───────────────────────────────────────────────────────────────

def _build_groq_llm(*, model_name: str, max_tokens: int = 512):
    from langchain_groq import ChatGroq

    load_dotenv(dotenv_path=_ENV_FILE, override=True)
    api_key = os.environ.get("GROQ_API_KEY", "").strip()
    if not api_key or api_key == "your_groq_api_key_here":
        raise ValueError("GROQ_API_KEY is not configured for multimodal ingestion.")
    return ChatGroq(
        api_key=api_key,
        model=model_name,
        temperature=0.2,
        max_tokens=max_tokens,
    )


def _summarize_text_block(text: str, *, instruction: str) -> str | None:
    if not settings.enable_summary_index or not text.strip():
        return None
    try:
        llm = _build_groq_llm(model_name=settings.groq_model, max_tokens=350)
        prompt = (
            f"{instruction}\n\n"
            "Keep the summary concise, retrieval-friendly, and factual.\n"
            "Respond with only the summary.\n\n"
            f"Content:\n{text[:settings.summary_max_chars]}"
        )
        response = llm.invoke(prompt)
        content = getattr(response, "content", "")
        return content.strip() if isinstance(content, str) and content.strip() else None
    except Exception as exc:
        logger.warning(f"Text summarization skipped: {exc}")
        return None


def _summarize_image(base64_image: str, *, mime_type: str, source: str, page: int | None) -> str | None:
    if not settings.enable_image_summaries or not base64_image.strip():
        return None
    try:
        llm = _build_groq_llm(model_name=settings.vision_model, max_tokens=450)
        prompt = (
            "Describe this document image for retrieval in a RAG system. "
            "Focus on what the image shows, visible labels/text, axes, legends, values, "
            "and why it matters in the document. Keep it concise but specific."
        )
        response = llm.invoke([
            HumanMessage(
                content=[
                    {"type": "text", "text": f"{prompt}\nSource: {source}\nPage: {page or 'unknown'}"},
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:{mime_type};base64,{base64_image}"},
                    },
                ]
            )
        ])
        content = getattr(response, "content", "")
        return content.strip() if isinstance(content, str) and content.strip() else None
    except Exception as exc:
        logger.warning(f"Image summarization skipped: {exc}")
        return None


# ── Loaders ────────────────────────────────────────────────────────────────────

def _document_id(file_path: str, user_id: int) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, f"{user_id}:{Path(file_path).name}"))


def _slugify(value: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", value.strip().lower()).strip("-")
    return slug or "section"


def _pdf_has_images(file_path: str) -> bool:
    try:
        import fitz

        with fitz.open(file_path) as pdf:
            for page in pdf:
                if page.get_images(full=True):
                    return True
    except Exception:
        return False
    return False


def _multimodal_assets_dir(file_path: str, document_id: str) -> Path:
    base = Path(file_path).parent / ".multimodal_assets" / document_id
    base.mkdir(parents=True, exist_ok=True)
    return base


def _pdf_assets_dir(file_path: str, document_id: str) -> Path:
    # Backward-compatible alias for existing PDF call sites.
    return _multimodal_assets_dir(file_path, document_id)


def remove_multimodal_assets(file_path: str, *, user_id: int) -> int:
    """
    Delete persisted multimodal extraction assets for a document.

    Returns 1 if an asset directory was removed, else 0.
    """
    document_id = _document_id(file_path, user_id)
    asset_dir = Path(file_path).parent / ".multimodal_assets" / document_id
    if asset_dir.exists():
        shutil.rmtree(asset_dir, ignore_errors=True)
        return 1
    return 0


def _load_pdf_pymupdf(file_path: str) -> list[Document]:
    """
    Load a PDF using PyMuPDF (fitz) for fast text extraction.
    Falls back to pypdf if fitz is unavailable.
    Returns one Document per page with page metadata.
    """
    try:
        import fitz  # pymupdf

        docs: list[Document] = []
        with fitz.open(file_path) as pdf:
            for page_num in range(len(pdf)):
                page = pdf[page_num]
                text = page.get_text("text")          # plain text extraction
                if not text.strip():
                    # Try harder with layout-preserving extraction
                    text = page.get_text("blocks")
                    if isinstance(text, list):
                        text = "\n".join(b[4] for b in text if isinstance(b[4], str))
                if text.strip():
                    docs.append(Document(
                        page_content=text.strip(),
                        metadata={
                            "source": Path(file_path).name,
                            "page": page_num + 1,
                            "total_pages": len(pdf),
                            "content_type": "text",
                            "extraction_method": "pymupdf",
                        },
                    ))
        logger.info(f"PyMuPDF: extracted text from {len(docs)}/{len(fitz.open(file_path))} pages")
        return docs

    except ImportError:
        logger.warning("PyMuPDF not installed, falling back to pypdf")
        from langchain_community.document_loaders import PyPDFLoader
        loader = PyPDFLoader(file_path)
        return loader.load()


def _load_pdf_multimodal(file_path: str, *, document_id: str, user_id: int) -> IngestionBundle:
    try:
        from unstructured.partition.pdf import partition_pdf
    except Exception as exc:
        logger.warning(f"Multimodal PDF loader unavailable, falling back to text-only path: {exc}")
        return IngestionBundle(raw_docs=_load_pdf_pymupdf(file_path), used_multimodal=False)

    filename = Path(file_path).name
    asset_dir = _pdf_assets_dir(file_path, document_id)
    raw_docs: list[Document] = []
    summary_docs: list[Document] = []
    image_count = 0

    try:
        chunks = partition_pdf(
            filename=file_path,
            infer_table_structure=True,
            strategy="hi_res",
            extract_image_block_types=["Image"],
            extract_image_block_to_payload=True,
            chunking_strategy="by_title",
            max_characters=4000,
            combine_text_under_n_chars=1500,
            new_after_n_chars=3000,
        )
    except Exception as exc:
        logger.warning(f"Multimodal PDF parsing failed, falling back to text-only path: {exc}")
        return IngestionBundle(raw_docs=_load_pdf_pymupdf(file_path), used_multimodal=False)

    for idx, chunk in enumerate(chunks):
        metadata = getattr(chunk, "metadata", None)
        page = getattr(metadata, "page_number", None)
        raw_text = (getattr(chunk, "text", "") or "").strip()
        chunk_type = type(chunk).__name__

        if raw_text:
            content_type = "table" if "Table" in chunk_type else "text"
            raw_docs.append(
                Document(
                    page_content=raw_text,
                    metadata={
                        "source": filename,
                        "page": page,
                        "content_type": content_type,
                        "section_title": getattr(metadata, "section", "") or "",
                        "extraction_method": "unstructured_hi_res",
                    },
                )
            )

            if content_type == "table":
                table_summary = _summarize_text_block(
                    raw_text,
                    instruction="Summarize this table for retrieval. Mention the main entities, metrics, and trends.",
                )
                if table_summary:
                    summary_docs.append(
                        Document(
                            page_content=table_summary,
                            metadata={
                                "source": filename,
                                "page": page,
                                "document_id": document_id,
                                "user_id": user_id,
                                "node_type": "table_summary",
                                "parent_id": document_id,
                                "table_index": idx,
                                "content_type": "table",
                            },
                        )
                    )

        image_base64 = getattr(metadata, "image_base64", None)
        image_mime_type = getattr(metadata, "image_mime_type", "image/jpeg") or "image/jpeg"
        if not image_base64 or image_count >= settings.vision_max_images_per_document:
            continue

        image_count += 1
        suffix = ".png" if "png" in image_mime_type else ".jpg"
        image_path = asset_dir / f"page-{page or 0}-image-{image_count}{suffix}"
        try:
            image_path.write_bytes(base64.b64decode(image_base64))
        except Exception as exc:
            logger.warning(f"Failed to persist extracted image for '{filename}': {exc}")
            image_path = asset_dir / f"missing-{image_count}{suffix}"

        image_summary = _summarize_image(
            image_base64,
            mime_type=image_mime_type,
            source=filename,
            page=page,
        )
        if image_summary:
            summary_docs.append(
                Document(
                    page_content=image_summary,
                    metadata={
                        "source": filename,
                        "page": page,
                        "document_id": document_id,
                        "user_id": user_id,
                        "node_type": "image_summary",
                        "parent_id": document_id,
                        "image_index": image_count,
                        "image_mime_type": image_mime_type,
                        "image_path": str(image_path),
                        "content_type": "image",
                    },
                )
            )

    return IngestionBundle(raw_docs=raw_docs, summary_docs=summary_docs, used_multimodal=True)


def _load_pdf(file_path: str, *, document_id: str, user_id: int) -> IngestionBundle:
    mode = settings.pdf_ingestion_mode.strip().lower()
    wants_multimodal = mode in {"multimodal", "hi_res", "rich"}

    if wants_multimodal or (settings.enable_image_summaries and _pdf_has_images(file_path)):
        bundle = _load_pdf_multimodal(file_path, document_id=document_id, user_id=user_id)
        if bundle.raw_docs:
            return bundle

    raw_docs = _load_pdf_pymupdf(file_path)
    if raw_docs:
        return IngestionBundle(raw_docs=raw_docs, used_multimodal=False)

    if settings.enable_ocr_fallback:
        return _load_pdf_multimodal(file_path, document_id=document_id, user_id=user_id)
    return IngestionBundle(raw_docs=[], used_multimodal=False)


def _load_docx_text(file_path: str) -> list[Document]:
    from langchain_community.document_loaders import Docx2txtLoader

    return Docx2txtLoader(file_path).load()


def _extract_docx_images(file_path: str) -> list[tuple[bytes, str, str]]:
    """
    Extract embedded raster images from a DOCX package.

    Returns tuples: (image_bytes, mime_type, suffix).
    """
    if Path(file_path).suffix.lower() != ".docx":
        return []

    supported_suffixes = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tif", ".tiff"}
    images: list[tuple[bytes, str, str]] = []
    try:
        with zipfile.ZipFile(file_path, "r") as archive:
            for member_name in archive.namelist():
                lowered = member_name.lower()
                if not lowered.startswith("word/media/"):
                    continue
                suffix = Path(lowered).suffix.lower()
                if suffix not in supported_suffixes:
                    continue
                payload = archive.read(member_name)
                if not payload:
                    continue
                mime_type = mimetypes.types_map.get(suffix, "image/jpeg")
                images.append((payload, mime_type, suffix))
    except Exception as exc:
        logger.warning(f"DOCX image extraction skipped for '{Path(file_path).name}': {exc}")
        return []
    return images


def _load_docx(file_path: str, *, document_id: str, user_id: int) -> IngestionBundle:
    filename = Path(file_path).name
    raw_docs = _load_docx_text(file_path)
    if not settings.enable_image_summaries:
        return IngestionBundle(raw_docs=raw_docs, used_multimodal=False)

    images = _extract_docx_images(file_path)
    if not images:
        return IngestionBundle(raw_docs=raw_docs, used_multimodal=False)

    asset_dir = _multimodal_assets_dir(file_path, document_id)
    summary_docs: list[Document] = []
    image_limit = max(0, settings.vision_max_images_per_document)
    for idx, (image_bytes, image_mime_type, suffix) in enumerate(images[:image_limit], start=1):
        image_path = asset_dir / f"docx-image-{idx}{suffix}"
        try:
            image_path.write_bytes(image_bytes)
        except Exception as exc:
            logger.warning(f"Failed to persist DOCX image for '{filename}': {exc}")
            image_path = asset_dir / f"missing-docx-image-{idx}{suffix}"

        image_summary = _summarize_image(
            base64.b64encode(image_bytes).decode("utf-8"),
            mime_type=image_mime_type,
            source=filename,
            page=None,
        )
        if image_summary:
            summary_docs.append(
                Document(
                    page_content=image_summary,
                    metadata={
                        "source": filename,
                        "page": None,
                        "document_id": document_id,
                        "user_id": user_id,
                        "node_type": "image_summary",
                        "parent_id": document_id,
                        "image_index": idx,
                        "image_mime_type": image_mime_type,
                        "image_path": str(image_path),
                        "content_type": "image",
                        "extraction_method": "docx_media",
                    },
                )
            )

    return IngestionBundle(raw_docs=raw_docs, summary_docs=summary_docs, used_multimodal=True)


def _load_text(file_path: str) -> list[Document]:
    from langchain_community.document_loaders import TextLoader
    return TextLoader(file_path, encoding="utf-8", autodetect_encoding=True).load()


def _tabular_rows_to_docs(
    all_rows: list[tuple | list],
    *,
    filename: str,
    metadata: dict | None = None,
) -> list[Document]:
    """Convert table-like rows into retrieval-friendly Documents."""
    if not all_rows:
        return []
    meta = {"source": filename}
    if metadata:
        meta.update(metadata)
    docs: list[Document] = []

    def _dedupe_headers(row_values: tuple | list) -> list[str]:
        headers: list[str] = []
        seen: dict[str, int] = {}
        for j, c in enumerate(row_values):
            raw = str(c).strip() if c is not None else ""
            base = re.sub(r"\s+", " ", raw) if raw else ""
            if not base:
                base = f"col{j + 1}"
            key = base.lower()
            seen[key] = seen.get(key, 0) + 1
            if seen[key] > 1:
                base = f"{base}_{seen[key]}"
            headers.append(base)
        return headers

    def _parse_date_like_string(value: str) -> date | None:
        clean = value.strip()
        if not clean:
            return None

        for fmt in (
            "%m/%d/%y",
            "%m/%d/%Y",
            "%Y-%m-%d",
            "%Y/%m/%d",
            "%d-%b-%Y",
            "%d %b %Y",
            "%b %d %Y",
            "%B %d %Y",
        ):
            try:
                return datetime.strptime(clean, fmt).date()
            except ValueError:
                continue

        # Common serialized datetime string from spreadsheets.
        try:
            return datetime.fromisoformat(clean.replace("Z", "+00:00")).date()
        except ValueError:
            return None

    def _date_variants(d: date) -> list[str]:
        return [
            f"{d.year}-{d.month:02d}-{d.day:02d}",
            f"{d.month}/{d.day}/{d.year}",
            f"{d.month}/{d.day}/{str(d.year)[-2:]}",
            d.strftime("%b %d %Y"),
            d.strftime("%B %d %Y"),
        ]

    def _format_cell_for_retrieval(value, *, column_name: str) -> str:
        if value is None:
            return ""

        if isinstance(value, datetime):
            d = value.date()
            variants = _date_variants(d)
            return f"{variants[0]} ({', '.join(variants[1:])})"
        if isinstance(value, date):
            variants = _date_variants(value)
            return f"{variants[0]} ({', '.join(variants[1:])})"

        if isinstance(value, bool):
            return "true" if value else "false"

        if isinstance(value, float):
            if value.is_integer():
                return str(int(value))
            return f"{value}".rstrip("0").rstrip(".")
        if isinstance(value, int):
            return str(value)

        text_value = str(value).strip()
        if not text_value:
            return ""

        parsed_date = _parse_date_like_string(text_value)
        if parsed_date:
            variants = _date_variants(parsed_date)
            extra = [v for v in variants if v != text_value]
            if extra:
                return f"{text_value} ({', '.join(extra)})"

        # If a date column came through as an untyped string/number, keep it
        # as-is so exact literal matching still works.
        if "date" in column_name.lower():
            return text_value
        return text_value

    def _canonical_cell_value(value: str) -> str:
        return value.split(" (", 1)[0].strip()

    # Try to detect header row (first row with multiple non-empty string cells).
    headers: list[str] | None = None
    data_start = 0
    for idx, row in enumerate(all_rows):
        str_cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
        if len(str_cells) >= 2:
            # Header rows are usually mostly non-numeric values.
            non_numeric = sum(1 for c in str_cells if not _is_number(c))
            if non_numeric >= len(str_cells) * 0.5:
                headers = _dedupe_headers(row)
                data_start = idx + 1
                break

    if headers:
        # Structured mode: produce "Column: Value" pairs per row.
        section_label = ""
        for row_index, row in enumerate(all_rows[data_start:], start=data_start + 1):
            cells = list(row)

            # Detect section headers (single non-empty cell spanning a row).
            non_empty = [(j, c) for j, c in enumerate(cells) if c is not None and str(c).strip()]
            if len(headers) > 1 and len(non_empty) == 1:
                section_label = str(non_empty[0][1]).strip()
                continue
            if not non_empty:
                continue

            row_pairs: list[tuple[str, str]] = []
            if section_label:
                row_pairs.append(("Section", section_label))
            for j, val in enumerate(cells):
                if val is None or not str(val).strip():
                    continue
                col_name = headers[j] if j < len(headers) else f"col{j}"
                row_pairs.append((col_name, _format_cell_for_retrieval(val, column_name=col_name)))

            if row_pairs:
                details = "\n".join(f"{col}: {val}" for col, val in row_pairs)
                compact = " | ".join(f"{col}={val}" for col, val in row_pairs)
                derived_facts: list[str] = []

                def _find_row_value(*needles: str, any_of: tuple[str, ...] | None = None) -> str:
                    for col, value in row_pairs:
                        col_l = col.lower()
                        if any_of and any(token in col_l for token in any_of):
                            return value
                        if needles and all(token in col_l for token in needles):
                            return value
                    return ""

                patient_name = _find_row_value(any_of=("patient name", "name", "patient"))
                medical_condition = _find_row_value(any_of=("medical condition", "condition", "diagnosis"))
                admission_date = _find_row_value("admission", "date", any_of=("admission date", "date of admission", "admitted"))
                discharge_date = _find_row_value("discharge", "date", any_of=("discharge date", "date of discharge", "discharged"))

                if medical_condition and admission_date:
                    if patient_name:
                        derived_facts.append(
                            "Patient "
                            f"{_canonical_cell_value(patient_name)} was admitted for "
                            f"{_canonical_cell_value(medical_condition)} on "
                            f"{_canonical_cell_value(admission_date)}."
                        )
                    else:
                        derived_facts.append(
                            f"Admission was for {_canonical_cell_value(medical_condition)} on "
                            f"{_canonical_cell_value(admission_date)}."
                        )

                if discharge_date:
                    if patient_name:
                        derived_facts.append(
                            f"Patient {_canonical_cell_value(patient_name)} was discharged on "
                            f"{_canonical_cell_value(discharge_date)}."
                        )
                    else:
                        derived_facts.append(
                            f"Discharge date is {_canonical_cell_value(discharge_date)}."
                        )

                text = (
                    f"Sheet: {meta.get('sheet', 'Sheet1')}\n"
                    f"Row: {row_index}\n"
                    f"{details}\n"
                    f"Record: {compact}"
                )
                if derived_facts:
                    text += "\nFacts: " + " ".join(derived_facts)
                row_meta = dict(meta)
                row_meta["row_number"] = row_index
                row_meta["content_type"] = "table_row"
                docs.append(Document(page_content=text, metadata=row_meta))
    else:
        # Fallback: one document per non-empty row for better retrieval precision.
        for row_index, row in enumerate(all_rows, start=1):
            vals = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if not vals:
                continue
            text = (
                f"Sheet: {meta.get('sheet', 'Sheet1')}\n"
                f"Row: {row_index}\n"
                f"Values: {' | '.join(vals)}"
            )
            row_meta = dict(meta)
            row_meta["row_number"] = row_index
            row_meta["content_type"] = "table_row"
            docs.append(Document(page_content=text, metadata=row_meta))

    return docs


def _load_excel(file_path: str) -> list[Document]:
    """
    Load an Excel file into readable, structured Documents.

    Strategy: read each sheet with openpyxl, detect header rows, and
    convert each data row into a "key: value" sentence so that both the
    embedding model and the LLM can understand the tabular relationships.
    """
    import openpyxl
    from openpyxl.utils.datetime import from_excel

    wb = openpyxl.load_workbook(file_path, data_only=True)
    filename = Path(file_path).name
    docs: list[Document] = []

    for sheet in wb.worksheets:
        sheet_rows: list[list] = []
        for row in sheet.iter_rows(values_only=False):
            values: list = []
            for cell in row:
                value = cell.value
                if value is None:
                    values.append(None)
                    continue
                if cell.is_date and isinstance(value, (int, float)):
                    try:
                        value = from_excel(value)
                    except Exception:
                        pass
                values.append(value)
            sheet_rows.append(values)
        docs.extend(
            _tabular_rows_to_docs(
                sheet_rows,
                filename=filename,
                metadata={"sheet": sheet.title},
            )
        )

    return docs


def _load_csv(file_path: str) -> list[Document]:
    """Load CSV files into structured Documents similar to Excel ingestion."""
    filename = Path(file_path).name
    rows: list[list[str]] = []
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            with open(file_path, "r", encoding=encoding, newline="") as f:
                sample = f.read(4096)
                f.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                rows = list(csv.reader(f, dialect))
            break
        except UnicodeDecodeError:
            continue
    return _tabular_rows_to_docs(rows, filename=filename, metadata={"sheet": "CSV"})


def _is_number(s: str) -> bool:
    try:
        float(s.replace(",", ""))
        return True
    except ValueError:
        return False


def _load_file(file_path: str, *, document_id: str, user_id: int) -> IngestionBundle:
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return _load_pdf(file_path, document_id=document_id, user_id=user_id)
    if ext in (".docx", ".doc"):
        return _load_docx(file_path, document_id=document_id, user_id=user_id)
    if ext in (".txt", ".md"):
        return IngestionBundle(raw_docs=_load_text(file_path))
    if ext in (".xlsx", ".xls"):
        return IngestionBundle(raw_docs=_load_excel(file_path))
    if ext == ".csv":
        return IngestionBundle(raw_docs=_load_csv(file_path))
    raise ValueError(f"Unsupported file extension: {ext}")


# ── Chunking ───────────────────────────────────────────────────────────────────

def _build_document_summary(
    *,
    source: str,
    file_path: str,
    document_id: str,
    user_id: int,
    raw_docs: list[Document],
) -> list[Document]:
    if not settings.enable_summary_index or not raw_docs:
        return []
    combined = "\n\n".join(doc.page_content for doc in raw_docs if doc.page_content.strip())
    summary = _summarize_text_block(
        combined,
        instruction="Summarize this document for retrieval. Mention the main topics, entities, and conclusions.",
    )
    if not summary:
        return []
    return [
        Document(
            page_content=summary,
            metadata={
                "source": source,
                "file_path": file_path,
                "document_id": document_id,
                "user_id": user_id,
                "node_type": "document_summary",
                "parent_id": document_id,
                "content_type": "summary",
            },
        )
    ]


def _build_section_summaries(
    *,
    source: str,
    file_path: str,
    document_id: str,
    user_id: int,
    raw_docs: list[Document],
) -> list[Document]:
    if (
        not settings.enable_summary_index
        or not settings.enable_section_summaries
        or not raw_docs
    ):
        return []

    grouped_sections: dict[str, list[Document]] = {}
    for doc in raw_docs:
        section_title = (doc.metadata or {}).get("section_title", "")
        if not isinstance(section_title, str):
            continue
        clean_title = section_title.strip()
        if not clean_title:
            continue
        grouped_sections.setdefault(clean_title, []).append(doc)

    if not grouped_sections:
        return []

    summaries: list[Document] = []
    for index, (section_title, section_docs) in enumerate(grouped_sections.items(), start=1):
        section_text = "\n\n".join(
            d.page_content for d in section_docs if d.page_content.strip()
        )
        if not section_text.strip():
            continue

        summary = _summarize_text_block(
            section_text,
            instruction=(
                f"Summarize this section titled '{section_title}' for retrieval. "
                "Mention key concepts, entities, and conclusions."
            ),
        )
        if not summary:
            continue

        pages = [
            (d.metadata or {}).get("page")
            for d in section_docs
            if isinstance((d.metadata or {}).get("page"), int)
        ]
        section_id = f"{document_id}:section:{index}:{_slugify(section_title)}"
        summaries.append(
            Document(
                page_content=summary,
                metadata={
                    "source": source,
                    "file_path": file_path,
                    "document_id": document_id,
                    "user_id": user_id,
                    "node_type": "section_summary",
                    "section_id": section_id,
                    "section_title": section_title,
                    "parent_id": document_id,
                    "content_type": "summary",
                    "page_start": min(pages) if pages else None,
                    "page_end": max(pages) if pages else None,
                },
            )
        )

    return summaries


def load_and_chunk(file_path: str, *, user_id: int) -> tuple[list[Document], list[Document], str]:
    """
    Load a file, split into overlapping chunks, and attach source metadata.
    Uses larger chunk sizes than the defaults so multi-page PDFs produce
    meaningful retrievable passages.
    """
    filename = Path(file_path).name
    document_id = _document_id(file_path, user_id)
    bundle = _load_file(file_path, document_id=document_id, user_id=user_id)
    raw_docs = bundle.raw_docs

    total_chars = sum(len(d.page_content) for d in raw_docs)
    logger.info(f"Loaded '{filename}': {len(raw_docs)} page/section(s), "
                f"{total_chars:,} chars total")

    if total_chars == 0:
        logger.warning(f"'{filename}' yielded no extractable text — "
                       "both text and multimodal extraction returned nothing.")
        return [], bundle.summary_docs, document_id

    # Use larger sizes for better semantic coherence
    chunk_size    = max(settings.chunk_size, 800)
    chunk_overlap = max(settings.chunk_overlap, 150)

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
        length_function=len,
        is_separator_regex=False,
    )

    chunks = splitter.split_documents(raw_docs)

    # Stamp every chunk with consistent source metadata
    for i, chunk in enumerate(chunks):
        chunk.metadata["source"]      = filename
        chunk.metadata["file_path"]   = file_path
        chunk.metadata["user_id"]     = user_id
        chunk.metadata["document_id"] = document_id
        chunk.metadata["node_type"]   = "raw_chunk"
        chunk.metadata["parent_id"]   = document_id
        chunk.metadata["chunk_index"] = i
        chunk.metadata["total_chunks"] = len(chunks)

    summary_docs = _build_document_summary(
        source=filename,
        file_path=file_path,
        document_id=document_id,
        user_id=user_id,
        raw_docs=raw_docs,
    )
    summary_docs.extend(
        _build_section_summaries(
            source=filename,
            file_path=file_path,
            document_id=document_id,
            user_id=user_id,
            raw_docs=raw_docs,
        )
    )
    summary_docs.extend(bundle.summary_docs)
    for i, doc in enumerate(summary_docs):
        doc.metadata.setdefault("source", filename)
        doc.metadata.setdefault("file_path", file_path)
        doc.metadata.setdefault("user_id", user_id)
        doc.metadata.setdefault("document_id", document_id)
        doc.metadata.setdefault("parent_id", document_id)
        doc.metadata.setdefault("summary_index", i)

    logger.info(f"'{filename}' → {len(chunks)} chunks "
                f"(size={chunk_size}, overlap={chunk_overlap})")
    return chunks, summary_docs, document_id


# ── Public API ─────────────────────────────────────────────────────────────────

def ingest_file(file_path: str, *, user_id: int) -> dict:
    """Ingest a single file into ChromaDB. Returns a summary dict."""
    filename = Path(file_path).name
    chunks, summary_docs, _document_id_value = load_and_chunk(file_path, user_id=user_id)
    if not chunks and not summary_docs:
        return {
            "file": filename,
            "chunks": 0,
            "ids": 0,
            "warning": "No content could be extracted. If this is a scanned PDF, "
                       "install the OCR-capable PDF dependencies and retry.",
        }

    # Replace any prior raw/summary nodes for this source to avoid duplicates.
    removed = delete_document_family(source=filename, user_id=user_id)
    if removed:
        logger.info(f"Removed {removed} existing indexed nodes for '{filename}' before re-ingest")

    raw_ids = add_raw_documents(chunks)
    summary_ids = add_summary_documents(summary_docs)
    warning = None
    if not chunks and summary_docs:
        warning = "Indexed summary-only content. Detailed raw-text answers may be limited for this document."
    return {
        "file": filename,
        "chunks": len(chunks),
        "ids": len(raw_ids) + len(summary_ids),
        "warning": warning,
    }


def ingest_directory(directory: str | None = None, *, user_id: int) -> list[dict]:
    """Walk the documents directory and ingest all supported files."""
    directory = directory or settings.documents_path
    results   = []

    for root, _, files in os.walk(directory):
        for fname in files:
            if Path(fname).suffix.lower() not in SUPPORTED_EXTENSIONS:
                continue
            fpath = os.path.join(root, fname)
            try:
                results.append(ingest_file(fpath, user_id=user_id))
            except Exception as exc:
                logger.error(f"Failed to ingest '{fname}': {exc}")
                results.append({"file": fname, "error": str(exc)})

    return results
